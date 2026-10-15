"""Material suplementario de la V21: el anexo A de la V20, renumerado como S1 a S7 (tablas y figuras con prefijo S).
Uso: python construir_suplemento_v21.py"""
import os, re, copy, unicodedata, zipfile, shutil
import docx
from docx.text.paragraph import Paragraph
from lxml import etree
ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
SRC = f'{ROOT}/01_MANUSCRITO/MANUSCRITO_QUITO_V20.docx'; OUT = f'{ROOT}/01_MANUSCRITO/MATERIAL_SUPLEMENTARIO_V21.docx'
WNS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'; W = lambda t: f'{{{WNS}}}{t}'
TITULO = 'Privación social y preparación ante inundaciones en Quito: una priorización parroquial con incertidumbre cuantificada'
d = docx.Document(SRC); paras = d.paragraphs; body = list(d.element.body.iterchildren()); sect = body[-1]
i_start = [i for i, p in enumerate(paras) if p.text.startswith('Anexo A Especificación')][0]
i_ref = [i for i, p in enumerate(paras) if p.text.strip() == 'Referencias'][0]
start_el, ref_el = paras[i_start]._p, paras[i_ref]._p
idx = {id(e): k for k, e in enumerate(body)}
keep = [e for e in body[idx[id(start_el)]:idx[id(ref_el)]]]
# fuera el bloque de la figura 8 (pasa a ser la figura 5 del texto principal)
cap8 = [p for p in paras if p.text.strip().startswith('Figura 8.')][0]._p
drop = set()
cur = cap8
for _ in range(3): drop.add(id(cur)); cur = cur.getnext()
keep = [e for e in keep if id(e) not in drop]
refs_all = [p for p in paras[i_ref + 1:] if p.text.strip()]
h1_tpl = [p for p in paras if p.text.strip() == 'Resumen'][0]; body_tpl = [p for p in paras if p.text.startswith('Los resultados muestran una mejora puntual')][0]
title_tpl = paras[0]; author_tpl = paras[1]
FIG = {'5': 'S1', '6': 'S2', '7': 'S3', '9': 'S4'}
TAB = {'5': 'S1', '6': 'S2', '7': 'S3', '8': 'S4', '9': 'S5', '10': 'S6', '11': 'S7'}
def remap(t):
    if t.startswith('Fuente: INEC (2024)'): return t            # tabla 1, tabla 3 y anexo 1 del documento del INEC
    t = re.sub(r'\bA\.(\d)\b', r'S\1', t)
    t = re.sub(r'\banexo S(\d)', r'sección S\1', t)
    t = t.replace('tablas 6 y 7', 'tablas S2 y S3')
    t = re.sub(r'\bfigura 8\b', 'figura 5 del texto principal', t)
    t = re.sub(r'\b(Figura|figura) (\d+)\b', lambda m: f'{m.group(1)} {FIG[m.group(2)]}' if m.group(2) in FIG else m.group(0), t)
    t = re.sub(r'\b(Tabla|tabla) (\d+)\b', lambda m: f'{m.group(1)} {TAB[m.group(2)]}' if m.group(2) in TAB else m.group(0), t)
    return t
def set_text(p, new):
    rs = p.runs
    if not rs: return
    rs[0].text = new
    for r in rs[1:]: r.text = ''
out = []
t = copy.deepcopy(title_tpl._p); pt = Paragraph(t, title_tpl._parent)
for r in pt.runs[1:]: r._r.getparent().remove(r._r)
pt.runs[0].text = 'Material suplementario'; out.append(t)
s1 = copy.deepcopy(author_tpl._p); ps1 = Paragraph(s1, author_tpl._parent)
for r in ps1.runs[1:]: r._r.getparent().remove(r._r)
ps1.runs[0].text = TITULO; out.append(s1)
a = copy.deepcopy(author_tpl._p); pa = Paragraph(a, author_tpl._parent)
for r in pa.runs[1:]: r._r.getparent().remove(r._r)
pa.runs[0].text = 'Wilson Gustavo Morales Lastra y Josseline Daniela Chávez Carpio'; out.append(a)
intro = copy.deepcopy(body_tpl._p); pi = Paragraph(intro, body_tpl._parent)
for r in pi.runs[1:]: r._r.getparent().remove(r._r)
pi.runs[0].text = ('Este material complementa el artículo. Reúne la especificación y el entrenamiento del modelo territorial (S1 a S4), su relación con la medición satelital de inundaciones (S5), '
                   'las pruebas de robustez adicionales (S6), la definición y verificación del NBI (S7) y las pruebas posestimación (S8). Las tablas y figuras se numeran con el prefijo S.')
out.append(intro)
for e in keep[1:]:                       # se omite el encabezado "Anexo A ..."
    e = copy.deepcopy(e)
    if e.tag == W('p') and not e.xpath('.//m:oMath'):
        p = Paragraph(e, None); txt = p.text
        if txt.strip():
            new = remap(txt)
            if new != txt: set_text(p, new)
    out.append(e)

# =================== S8: pruebas posestimación ===================
import pandas as pd
R09 = f'{ROOT}/09_ANALISIS_ROBUSTEZ'; FIGDIR = f'{ROOT}/10_FIGURAS'
h2_tpl = [p for p in paras if p.text.startswith('A.7 Definición y verificación del NBI')][0]
cap_tpl = [p for p in paras if p.text.startswith('Tabla 5. Particiones de evaluación')][0]
note_tpl = [p for p in paras if p.text.startswith('Fuente: particiones congeladas')][0]
tbl_tpl = [t for t in d.tables if t.rows[0].cells[0].text.strip() == 'Protocolo'][0]._tbl
ref_tpl = [p for p in paras if p.text.startswith('Elfwing, S.')][0]
NB = ' '
def nbs(t): return re.sub(r'(?<=\d) (?=\d{3}(?!\d))', NB, t)
def cl(tpl, text):
    new = copy.deepcopy(tpl._p)
    for el in new.xpath('.//w:lastRenderedPageBreak'): el.getparent().remove(el)
    p = Paragraph(new, tpl._parent)
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    p.runs[0].text = nbs(text); return new
def make_table(widths, header, rows):
    t = copy.deepcopy(tbl_tpl); trs = t.xpath('./w:tr'); hdr, dat = copy.deepcopy(trs[0]), copy.deepcopy(trs[1])
    for tr in trs: t.remove(tr)
    grid = t.find(W('tblGrid'))
    for g in list(grid): grid.remove(g)
    for w in widths: etree.SubElement(grid, W('gridCol')).set(W('w'), str(w))
    def build(tpl_tr, vals):
        tr = copy.deepcopy(tpl_tr); tcs = tr.xpath('./w:tc'); tc0 = copy.deepcopy(tcs[0])
        for tc in tcs: tr.remove(tc)
        for w, v in zip(widths, vals):
            tc = copy.deepcopy(tc0); tc.find(W('tcPr')).find(W('tcW')).set(W('w'), str(w))
            ps = tc.xpath('./w:p')
            for ex in ps[1:]: tc.remove(ex)
            p = Paragraph(ps[0], None); rs = p.runs
            for rr in rs[1:]: rr._r.getparent().remove(rr._r)
            rs[0].text = nbs(str(v)); tr.append(tc)
        return tr
    t.append(build(hdr, header))
    for row in rows: t.append(build(dat, row))
    return t
def picture(path, width_cm=16):
    from docx.shared import Cm
    new = copy.deepcopy(cap_tpl._p)
    for r in new.xpath('./w:r'): new.remove(r)
    Paragraph(new, cap_tpl._parent).add_run().add_picture(path, width=Cm(width_cm)); return new
co = pd.read_csv(f'{R09}/A8_coeficientes_vif.csv'); inf = pd.read_csv(f'{R09}/A8_influencia_bloques.csv'); loc = pd.read_csv(f'{R09}/A8_incertidumbre_localizacion.csv')
hp = pd.read_csv(f'{R09}/A8_sensibilidad_hiperparametros.csv'); an = pd.read_csv(f'{R09}/A8_coe_por_anio.csv'); dep = pd.read_csv(f'{R09}/A8_dependencia_bloques.csv')
nc = co[~co.constante].copy(); n_exc = int(nc.ic_excluye_cero.sum()); vif10 = int((nc.vif_fondo > 10).sum())
ht, he = hp[hp.protocolo == 'temporal'], hp[hp.protocolo == 'espacial']
loc = loc.set_index('posiciones'); l25, lc, l9 = loc.iloc[0], loc.iloc[1], loc.iloc[2]
cb = co.set_index('variable')
S8 = []
S8.append(cl(h2_tpl, 'S8 Pruebas posestimación'))
S8.append(cl(body_tpl, 'Esta sección reúne las pruebas posestimación del ancla loglineal de la variante aplicada (temporal, con vías y FABDEM, con penalización 0.1 y temperado 1.25), que concentra casi todo el desempeño del modelo (sección S6). Los intervalos usan bootstrap por bloques de 5 km con 300 réplicas (Efron, 1979; Künsch, 1989).'))
S8.append(cl(body_tpl, f"Colinealidad. Las 22 covariables no constantes tienen una matriz de correlaciones de rango 21: las clases de cobertura suman uno por construcción y quedan identificadas por la penalización L2. Con la clase más frecuente (cobertura 10) como referencia, el número de condición de la matriz de correlaciones es 131 en el fondo y 83 en las posiciones de los reportes. Dos covariables superan un factor de inflación de la varianza de 10 (construido a 100 m, {cb.loc['construido_100m','vif_fondo']:.1f}, y a 500 m, {cb.loc['construido_500m','vif_fondo']:.1f}; O'Brien, 2007; Dormann et al., 2013), lo que explica que el coeficiente a 500 m cambie de signo en {100 * (1 - cb.loc['construido_500m','estabilidad_signo']):.0f} % de las réplicas y no deba interpretarse por separado (tabla S8)."))
S8.append(cl(body_tpl, f"Estabilidad de los coeficientes. De las 22 covariables no constantes, cinco tienen un intervalo del 95 % que excluye cero: el relieve local ({cb.loc['relieve_local_m','coeficiente_publicado']:.2f}), el construido a 100 m ({cb.loc['construido_100m','coeficiente_publicado']:.2f}) y la distancia a vías ({cb.loc['log_distancia_vial','coeficiente_publicado']:.2f}), que conservan su signo en al menos 99.7 % de las réplicas, y dos clases de cobertura raras (agua permanente y humedal) con coeficientes cercanos a cero. Las demás covariables conservan el signo en 56 % a 93 % de las réplicas. El desempeño descansa, por tanto, en tres covariables, en línea con la equivalencia del modelo de tres variables de la tabla S4."))
rows8 = []
lab = lambda v: v.replace('_', ' ')
for _, r in nc.iterrows():
    rows8.append([lab(r.variable), f'{r.coeficiente_publicado:.2f}', f'{r.ic95_inf:.2f} a {r.ic95_sup:.2f}', f'{100 * r.estabilidad_signo:.0f}', ('ref.' if r.clase_referencia_vif else f'{r.vif_fondo:.1f}')])
S8.append(cl(cap_tpl, 'Tabla S8. Coeficientes del ancla, estabilidad del signo y colinealidad'))
S8.append(make_table([3300, 1300, 2100, 1172, 1200], ['Covariable', 'Coeficiente', 'IC 95 %', 'Signo estable (%)', 'VIF'], rows8))
S8.append(cl(note_tpl, 'Fuente: ancla reajustada con 300 réplicas por bloques de 5 km. Elaboración propia. Coeficientes sobre covariables estandarizadas. VIF calculado en el fondo con cobertura 10 como clase de referencia. Se omiten tres clases de cobertura constantes (70, 95 y 100).'))
S8.append(cl(body_tpl, f"Influencia. Al excluir cada uno de los 42 bloques de entrenamiento (jackknife; Belsley et al., 1980), el coeficiente que más cambia lo hace en {inf.max_dif_estandarizada.max():.2f} errores estándar y el desempeño de prueba varía entre {inf.dif_logdens_test.min():.3f} y +{inf.dif_logdens_test.max():.3f} de log densidad y entre {100 * inf.dif_captura_test.min():.1f} y +{100 * inf.dif_captura_test.max():.1f} puntos de captura: ningún bloque determina el resultado."))
S8.append(cl(body_tpl, f"Incertidumbre de localización. Entrenar solo con la posición más central de cada reporte (la de menor distancia media a las demás en el espacio de covariables) y evaluar con las 25 posiciones da {lc.logdens_prueba_25_posiciones:.3f} de log densidad y {lc.captura_25_posiciones:.3f} de captura, frente a {l25.logdens_prueba_25_posiciones:.3f} y {l25.captura_25_posiciones:.3f} con las 25; con las nueve posiciones más centrales, {l9.logdens_prueba_25_posiciones:.3f} y {l9.captura_25_posiciones:.3f}. Evaluar solo en la posición central sobrestima el desempeño ({lc.logdens_prueba_mismas_posiciones:.3f} y {lc.captura_mismas_posiciones:.3f}). Los coeficientes cambian como máximo {lc.max_dif_coef_vs_25:.2f}."))
mej_t = ht.loc[ht.logdens_prueba.idxmax()]; mej_e = he.loc[he.logdens_prueba.idxmax()]; pub_e = he[he.publicado].iloc[0]; pub_t = ht[ht.publicado].iloc[0]
S8.append(cl(body_tpl, f"Penalización y temperado. En las 20 combinaciones de la grilla, la log densidad de prueba temporal va de {ht.logdens_prueba.min():.3f} a {ht.logdens_prueba.max():.3f} (publicada: {pub_t.logdens_prueba:.3f}) y la espacial de {he.logdens_prueba.min():.3f} a {he.logdens_prueba.max():.3f} (publicada: {pub_e.logdens_prueba:.3f}). El temperado domina: en la prueba temporal el promedio sube de {ht[ht.temperado == 0.25].logdens_prueba.mean():.2f} con 0.25 a {ht[ht.temperado == 1.25].logdens_prueba.mean():.2f} con 1.25, y en la espacial es máximo con 0.75 ({he[he.temperado == 0.75].logdens_prueba.mean():.2f}). El valor publicado queda a {mej_t.logdens_prueba - pub_t.logdens_prueba:.2f} del mejor de la prueba temporal y a {mej_e.logdens_prueba - pub_e.logdens_prueba:.2f} del mejor de la espacial, de modo que la selección por calibración es razonable, pero no óptima fuera de la zona de entrenamiento."))
rho_txt = '−0.46 (p = 0.08)'
S8.append(cl(body_tpl, f"Estabilidad temporal. Con los 2 257 eventos del COE de 2010 a mayo de 2024, la captura anual del 10 % superior del construido está entre {100 * an.captura.min():.1f} % y {100 * an.captura.max():.1f} %, y el límite inferior del intervalo de Wilson (1927) supera 10 % en los {len(an)} años (figura S5b). La correlación de rangos entre año y captura es de {rho_txt}: no hay tendencia significativa."))
S8.append(cl(body_tpl, f"Dependencia espacial. Los eventos del COE presentan una correlación intraclase de {dep.icc.iloc[0]:.2f} entre bloques de 5 km en 2022 a mayo de 2024 y de {dep.icc.iloc[1]:.2f} en 2010 a 2024, con efectos de diseño de {dep.efecto_diseno.iloc[0]:.1f} y {dep.efecto_diseno.iloc[1]:.1f} (Kish, 1965); el tamaño efectivo es de {dep.n_efectivo.iloc[0]:.0f} de {int(dep.n.iloc[0])} y de {dep.n_efectivo.iloc[1]:.0f} de {int(dep.n.iloc[1])} eventos. Los intervalos por bloques de las tablas del texto incorporan esa dependencia."))
S8.append(cl(cap_tpl, 'Figura S5. Coeficientes del ancla y estabilidad anual de la captura'))
S8.append(picture(f'{FIGDIR}/figS5_posestimacion.png'))
S8.append(cl(note_tpl, 'Fuente: ancla reajustada con bootstrap por bloques de 5 km y eventos de inundación del COE municipal. Elaboración propia. Panel a: azul, intervalo que excluye cero; gris, intervalo que incluye cero. Panel b: intervalos de Wilson del 95 %; la línea punteada es el azar (10 %).'))
S8.append(cl(cap_tpl, 'Tabla S9. Resumen de las pruebas posestimación'))
S8.append(make_table([2300, 3900, 2872], ['Prueba', 'Estadístico', 'Lectura'], [
 ['Colinealidad', f'Rango 21 de 22; condición 131 (fondo) y 83 (reportes); VIF > 10 en {vif10} covariables', 'Dos covariables de construido son colineales'],
 ['Estabilidad de coeficientes', f'{n_exc} de 22 con IC sin cero; tres estables en signo (≥ 99.7 %)', 'El desempeño descansa en tres covariables'],
 ['Jackknife por bloque', f'42 bloques; máximo {inf.max_dif_estandarizada.max():.2f} errores estándar; Δ log densidad {inf.dif_logdens_test.min():.3f} a +{inf.dif_logdens_test.max():.3f}', 'Ningún bloque determina el resultado'],
 ['Localización', f'Posición central: {lc.logdens_prueba_25_posiciones:.3f} y {100 * lc.captura_25_posiciones:.1f} %; nueve posiciones: {l9.logdens_prueba_25_posiciones:.3f} y {100 * l9.captura_25_posiciones:.1f} %', 'Resultado estable; evaluar solo en la posición central lo sobrestima'],
 ['Penalización y temperado', f'Log densidad {ht.logdens_prueba.min():.2f} a {ht.logdens_prueba.max():.2f} (temporal) y {he.logdens_prueba.min():.2f} a {he.logdens_prueba.max():.2f} (espacial)', 'El temperado domina; el valor publicado es razonable'],
 ['Estabilidad por año (COE)', f'Captura {100 * an.captura.min():.1f} % a {100 * an.captura.max():.1f} %; IC inferior > 10 % en {len(an)} años', 'Sin tendencia (Spearman −0.46; p = 0.08)'],
 ['Dependencia por bloques', f'ICC {dep.icc.iloc[1]:.2f} a {dep.icc.iloc[0]:.2f}; efecto de diseño {dep.efecto_diseno.iloc[0]:.1f} a {dep.efecto_diseno.iloc[1]:.1f}', 'Justifica el bootstrap por bloques']]))
S8.append(cl(note_tpl, 'Fuente: pruebas posestimación del estudio. Elaboración propia. Detalle en A8_*.csv de la carpeta de análisis de robustez.'))
out.extend(S8)
NEW_REF_TEXTS = [
 'Belsley, D. A., Kuh, E., & Welsch, R. E. (1980). Regression diagnostics: Identifying influential data and sources of collinearity. Wiley.',
 'Dormann, C. F., Elith, J., Bacher, S., Buchmann, C., Carl, G., Carré, G., García Marquéz, J. R., Gruber, B., Lafourcade, B., Leitão, P. J., Münkemüller, T., McClean, C., Osborne, P. E., Reineking, B., Schröder, B., Skidmore, A. K., Zurell, D., & Lautenbach, S. (2013). Collinearity: A review of methods to deal with it and a simulation study evaluating their performance. Ecography, 36(1), 27-46. https://doi.org/10.1111/j.1600-0587.2012.07348.x',
 'Kish, L. (1965). Survey sampling. Wiley.',
 "O'Brien, R. M. (2007). A caution regarding rules of thumb for variance inflation factors. Quality & Quantity, 41(5), 673-690. https://doi.org/10.1007/s11135-006-9018-6",
 'Wilson, E. B. (1927). Probable inference, the law of succession, and statistical inference. Journal of the American Statistical Association, 22(158), 209-212. https://doi.org/10.1080/01621459.1927.10502953']
for _t in NEW_REF_TEXTS: refs_all.append(Paragraph(cl(ref_tpl, _t), ref_tpl._parent))
# referencias citadas en el suplemento
ALIAS = {'Instituto Nacional de Estadística y Censos': 'INEC', 'Servicio Nacional de Gestión de Riesgos': 'SNGR'}
def ref_key(t):
    m = re.match(r'^(.+?)\. \((\d{4}|s\. f\.(?:-[a-z])?)\)', t)
    if not m: return None
    first = m.group(1).split(',')[0].strip(); return ALIAS.get(first, first), m.group(2)
text = '\n'.join(''.join((x.text or '') for x in e.iter(W('t'))) for e in out)
h = copy.deepcopy(h1_tpl._p); ph = Paragraph(h, h1_tpl._parent)
for r in ph.runs[1:]: r._r.getparent().remove(r._r)
ph.runs[0].text = 'Referencias'; out.append(h)
n = 0
for p in refs_all:
    k = ref_key(p.text)
    if k and re.search(re.escape(k[0]) + r'[^;)]{0,70}?' + re.escape(k[1]), text): out.append(copy.deepcopy(p._p)); n += 1
for e in body[:-1]: d.element.body.remove(e)
for e in out: sect.addprevious(e)
doc_xml = etree.tostring(d.element).decode()
for rid, rel in list(d.part.rels.items()):
    if 'image' in rel.reltype and f'"{rid}"' not in doc_xml: d.part.drop_rel(rid)
d.save(OUT)
tmp = OUT + '.tmp'
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename == 'word/footnotes.xml': data = re.sub(r'<w:footnote [^>]*w:id="1"[^>]*>.*?</w:footnote>', '', data.decode('utf-8'), flags=re.S).encode('utf-8')
        zout.writestr(it, data)
shutil.move(tmp, OUT)
print('suplemento guardado | referencias:', n)
