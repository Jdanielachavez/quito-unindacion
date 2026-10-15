"""Estilo de revista: tipografía serif (Times New Roman), título y autoría centrados, encabezados en negro,
y reescritura de la introducción (apertura, brecha, objetivo). Edita en su lugar los .docx de 00_ENTREGA_FINAL."""
import sys, copy, docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
D = '/Users/daniela/Desktop/PAPER INUNDACION/00_ENTREGA_FINAL/'
FONT = 'Times New Roman'
def fuentes(doc):
    st = doc.styles.element
    for rf in st.iter(qn('w:rFonts')):
        for a in list(rf.attrib):
            if 'Theme' in a or a.endswith('}cstheme'): del rf.attrib[a]
        for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'): rf.set(qn('w:' + a), FONT)
    for s in doc.styles:
        try:
            if s.font is not None and s.type in (1, 2, 3):
                s.font.name = FONT
                rf = s.element.get_or_add_rPr().get_or_add_rFonts()
                for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'): rf.set(qn('w:' + a), FONT)
                for a in list(rf.attrib):
                    if 'Theme' in a or a.endswith('}cstheme'): del rf.attrib[a]
        except Exception: pass
    # fuentes directas en el cuerpo (excepto matemáticas)
    for rf in doc.element.body.iter(qn('w:rFonts')):
        if rf.get(qn('w:ascii')) in ('Cambria Math',): continue
        for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'): rf.set(qn('w:' + a), FONT)
        for a in list(rf.attrib):
            if 'Theme' in a: del rf.attrib[a]
    for n in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        if n in [s.name for s in doc.styles]:
            f = doc.styles[n].font; f.color.rgb = RGBColor(0, 0, 0); f.bold = True; f.name = FONT
    doc.styles['Title'].font.size = Pt(16); doc.styles['Heading 1'].font.size = Pt(13); doc.styles['Heading 2'].font.size = Pt(11.5)
    doc.styles['Normal'].font.size = Pt(11)
def titulo(doc):
    P = doc.paragraphs
    P[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for p in P[1:5]:
        if p.text.startswith(('Wilson', 'Códigos JEL', 'Palabras clave')): p.alignment = WD_ALIGN_PARAGRAPH.CENTER if p.text.startswith('Wilson') else WD_ALIGN_PARAGRAPH.JUSTIFY
INTRO = {
 'Las inundaciones urbanas no afectan por igual': 'Las inundaciones urbanas no afectan por igual a todos los hogares. La exposición física determina dónde se acumula el agua, pero la capacidad de prepararse, responder y recuperarse depende de las condiciones sociales y de los servicios disponibles (Cutter et al., 2003). La evidencia global indica que integrar exposición y vulnerabilidad modifica tanto la magnitud del riesgo como la distribución de las medidas de preparación (Rentschler et al., 2022; Fox et al., 2024; Mester et al., 2025). En consecuencia, decidir dónde prepararse exige combinar, y no solo superponer, una medida de dónde es probable que se acumule el agua con una medida de quién enfrenta las mayores carencias.',
 'Quito ilustra el problema de asignación': 'Quito ilustra este problema de asignación. Su relieve, su urbanización y el acceso a servicios varían entre las 65 parroquias del Distrito Metropolitano (Watson et al., 2022), de modo que una parroquia con amplia superficie susceptible puede encabezar una selección por extensión, mientras que otra con menor superficie y mayores carencias adquiere prioridad si el objetivo es atender necesidades básicas. La evidencia económica ecuatoriana refuerza la necesidad de considerar las condiciones de respuesta además de la exposición: las respuestas a los choques de lluvia difieren por actividad y condición laboral (Canavire-Bacarreza et al., 2026), se acumulan con la recurrencia (Llerena Pinto et al., 2025) y pueden extenderse al capital humano (Rosales-Rueda, 2018). Con frecuencia, sin embargo, la elección del criterio se toma sin hacerla explícita y sin cuantificar cuánto depende de los supuestos.',
 'La literatura aplicada presenta tres limitaciones': 'La literatura aplicada presenta tres limitaciones recurrentes. Primero, los puntajes de susceptibilidad construidos con reportes de eventos suelen interpretarse como probabilidades de inundación, aunque estiman la concentración relativa de reportes (Warton & Shepherd, 2010) y están afectados por el sesgo de registro. Segundo, las tasas sociales provienen de capas agregadas cuyo cálculo rara vez se verifica contra las cifras oficiales. Tercero, la estabilidad de la priorización se examina frente a pocos supuestos y casi nunca frente a la incertidumbre muestral de los reportes.',
 'La evidencia económica ecuatoriana sugiere': 'La evidencia económica ecuatoriana citada en la introducción (Canavire-Bacarreza et al., 2026; Llerena Pinto et al., 2025; Rosales-Rueda, 2018) sugiere por qué las prioridades deben considerar condiciones de respuesta y no solo exposición. Nuestro cruce censal no observa ingresos ni composición laboral dentro del decil, por lo que esa vía queda como motivación para vincular observaciones de hogares en episodios nuevos y no como estimación de pérdidas.',
}
if __name__ == '__main__':
    for f, main in (('1_MANUSCRITO_V21.docx', True), ('2_MATERIAL_SUPLEMENTARIO_V21.docx', False)):
        d = docx.Document(D + f); fuentes(d); titulo(d)
        if main:
            hecho = 0
            for p in d.paragraphs:
                for k, v in INTRO.items():
                    if p.text.startswith(k):
                        p.runs[0].text = v
                        for r in p.runs[1:]: r._r.getparent().remove(r._r)
                        hecho += 1
            assert hecho == 4, hecho
        d.save(D + f); print('ok', f)
