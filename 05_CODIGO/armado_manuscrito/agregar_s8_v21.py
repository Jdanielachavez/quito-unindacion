"""Agrega la sección S8 (pruebas posestimación) a MATERIAL_SUPLEMENTARIO_V21.docx y sus
referidos en MANUSCRITO_QUITO_V21.docx, editando los .docx vigentes en su lugar
(V20, la fuente de construir_v21.py, ya no está en la carpeta)."""
import copy, re, shutil, sys
import docx, pandas as pd
from docx.text.paragraph import Paragraph
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = '/Users/daniela/Desktop/PAPER INUNDACION'
MAN, SUP = f'{ROOT}/01_MANUSCRITO/MANUSCRITO_QUITO_V21.docx', f'{ROOT}/01_MANUSCRITO/MATERIAL_SUPLEMENTARIO_V21.docx'
R09, FIG = f'{ROOT}/09_ANALISIS_ROBUSTEZ', f'{ROOT}/10_FIGURAS'
NB = ' '
def nbs(t): return re.sub(r'(?<=\d) (?=\d{3}(?!\d))', NB, t)

co = pd.read_csv(f'{R09}/A8_coeficientes_vif.csv'); inf = pd.read_csv(f'{R09}/A8_influencia_bloques.csv')
loc = pd.read_csv(f'{R09}/A8_incertidumbre_localizacion.csv'); hp = pd.read_csv(f'{R09}/A8_sensibilidad_hiperparametros.csv')
an = pd.read_csv(f'{R09}/A8_coe_por_anio.csv'); dep = pd.read_csv(f'{R09}/A8_dependencia_bloques.csv')
nc = co[~co.constante]; cb = co.set_index('variable')
n_exc, vif10 = int(nc.ic_excluye_cero.sum()), int((nc.vif_fondo > 10).sum())
assert len(nc) == 22 and n_exc == 5 and vif10 == 2
ht, he = hp[hp.protocolo == 'temporal'], hp[hp.protocolo == 'espacial']
pub_t, pub_e = ht[ht.publicado].iloc[0], he[he.publicado].iloc[0]
mej_t, mej_e = ht.loc[ht.logdens_prueba.idxmax()], he.loc[he.logdens_prueba.idxmax()]
l25, lc, l9 = loc.iloc[0], loc.iloc[1], loc.iloc[2]
from scipy.stats import spearmanr
rho, pv = spearmanr(an.anio, an.captura)
assert round(rho, 2) == -0.46

# ---------------- suplemento ----------------
shutil.copy(SUP, f'{SUP}.bak')
d = docx.Document(SUP)
P = d.paragraphs
assert not any(p.text.startswith('S8 ') for p in P)
find = lambda s: [p for p in P if p.text.startswith(s)][0]
h2_tpl, body_tpl = find('S7 Definición'), find('El recálculo coincide')
cap_tpl, note_tpl = find('Tabla S7.'), find('Fuente: INEC (2024), tabla 1')
tbl_tpl = d.tables[-1]._tbl
ref_head = find('Referencias')

def cl(tpl, text):
    new = copy.deepcopy(tpl._p)
    for el in new.xpath('.//w:lastRenderedPageBreak'): el.getparent().remove(el)
    p = Paragraph(new, tpl._parent)
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    p.runs[0].text = nbs(text); return new

def make_table(widths, header, rows):
    t = copy.deepcopy(tbl_tpl); trs = t.xpath('./w:tr'); hdr, dat = copy.deepcopy(trs[0]), copy.deepcopy(trs[1])
    for tr in trs: t.remove(tr)
    grid = t.find(qn('w:tblGrid'))
    for g in list(grid): grid.remove(g)
    from lxml import etree
    for w in widths: etree.SubElement(grid, qn('w:gridCol')).set(qn('w:w'), str(w))
    def build(tpl_tr, vals):
        tr = copy.deepcopy(tpl_tr); tcs = tr.xpath('./w:tc'); tc0 = copy.deepcopy(tcs[0])
        for tc in tcs: tr.remove(tc)
        for w, v in zip(widths, vals):
            tc = copy.deepcopy(tc0); tc.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'), str(w))
            ps = tc.xpath('./w:p')
            for ex in ps[1:]: tc.remove(ex)
            p = Paragraph(ps[0], None)
            for rr in p.runs[1:]: rr._r.getparent().remove(rr._r)
            p.runs[0].text = nbs(str(v)); tr.append(tc)
        return tr
    t.append(build(hdr, header))
    for row in rows: t.append(build(dat, row))
    return t

def picture(path, width_cm=16):
    new = copy.deepcopy(body_tpl._p)
    for r in new.xpath('./w:r'): new.remove(r)
    p = Paragraph(new, body_tpl._parent); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm)); return new

S8 = [cl(h2_tpl, 'S8 Pruebas posestimación')]
add = lambda t: S8.append(cl(body_tpl, t))
add('Esta sección reúne las pruebas posestimación del ancla loglineal de la variante aplicada (temporal, con vías y FABDEM; penalización 0.1 y temperado 1.25), que concentra casi todo el desempeño del modelo (sección S6). Los intervalos usan bootstrap por bloques de 5 km con 300 réplicas (Efron, 1979; Künsch, 1989).')
add(f"Colinealidad. Las 22 covariables no constantes tienen una matriz de correlaciones de rango 21: las clases de cobertura suman uno por construcción y quedan identificadas por la penalización L2. Con la clase más frecuente (cobertura 10) como referencia, el número de condición de la matriz de correlaciones es 131 en el fondo y 83 en las posiciones de los reportes. Dos covariables superan un factor de inflación de la varianza de 10 (construido a 100 m, {cb.loc['construido_100m','vif_fondo']:.1f}, y a 500 m, {cb.loc['construido_500m','vif_fondo']:.1f}; O'Brien, 2007; Dormann et al., 2013), lo que explica que el coeficiente a 500 m cambie de signo en {100 * (1 - cb.loc['construido_500m','estabilidad_signo']):.0f} % de las réplicas y no deba interpretarse por separado (tabla S8).")
add(f"Estabilidad de los coeficientes. De las 22 covariables no constantes, cinco tienen un intervalo del 95 % que excluye cero: el relieve local ({cb.loc['relieve_local_m','coeficiente_publicado']:.2f}), el construido a 100 m (+{cb.loc['construido_100m','coeficiente_publicado']:.2f}) y la distancia a vías ({cb.loc['log_distancia_vial','coeficiente_publicado']:.2f}), que conservan su signo en al menos 99.7 % de las réplicas, y dos clases de cobertura raras (cobertura 80 y 90) con coeficientes cercanos a cero. Las demás covariables conservan el signo en 56 % a 93 % de las réplicas. El desempeño descansa, por tanto, en tres covariables, en línea con la equivalencia del modelo de tres variables de la tabla S4.")
rows8 = [[r.variable.replace('_', ' '), f'{r.coeficiente_publicado:.2f}', f'{r.ic95_inf:.2f} a {r.ic95_sup:.2f}', f'{100 * r.estabilidad_signo:.0f}', ('ref.' if r.clase_referencia_vif else f'{r.vif_fondo:.1f}')] for _, r in nc.iterrows()]
S8.append(cl(cap_tpl, 'Tabla S8. Coeficientes del ancla, estabilidad del signo y colinealidad'))
S8.append(make_table([3300, 1300, 2100, 1172, 1200], ['Covariable', 'Coeficiente', 'IC 95 %', 'Signo estable (%)', 'VIF'], rows8))
S8.append(cl(note_tpl, 'Fuente: ancla reajustada con 300 réplicas por bloques de 5 km. Elaboración propia. Coeficientes sobre covariables estandarizadas. VIF calculado en el fondo con cobertura 10 como clase de referencia (ref.). Se omiten las covariables constantes en el soporte.'))
add(f"Influencia. Al excluir cada uno de los 42 bloques de entrenamiento (jackknife; Belsley et al., 1980), el coeficiente que más cambia lo hace en {inf.max_dif_estandarizada.max():.2f} errores estándar y el desempeño de prueba varía entre {inf.dif_logdens_test.min():.3f} y +{inf.dif_logdens_test.max():.3f} de log densidad y entre {100 * inf.dif_captura_test.min():.1f} y +{100 * inf.dif_captura_test.max():.1f} puntos de captura: ningún bloque determina el resultado.")
add(f"Incertidumbre de localización. Entrenar solo con la posición más central de cada reporte y evaluar con las 25 posiciones da {lc.logdens_prueba_25_posiciones:.3f} de log densidad y {100 * lc.captura_25_posiciones:.1f} % de captura, frente a {l25.logdens_prueba_25_posiciones:.3f} y {100 * l25.captura_25_posiciones:.1f} % con las 25; con las nueve posiciones más centrales, {l9.logdens_prueba_25_posiciones:.3f} y {100 * l9.captura_25_posiciones:.1f} %. Evaluar solo en la posición central sobrestima el desempeño ({lc.logdens_prueba_mismas_posiciones:.3f} y {100 * lc.captura_mismas_posiciones:.1f} %). Los coeficientes cambian como máximo {lc.max_dif_coef_vs_25:.2f}.")
add(f"Penalización y temperado. En las 20 combinaciones de la grilla, la log densidad de prueba temporal va de {ht.logdens_prueba.min():.3f} a {ht.logdens_prueba.max():.3f} (publicada: {pub_t.logdens_prueba:.3f}) y la espacial de {he.logdens_prueba.min():.3f} a {he.logdens_prueba.max():.3f} (publicada: {pub_e.logdens_prueba:.3f}). El temperado domina: en la prueba temporal el promedio sube de {ht[ht.temperado == 0.25].logdens_prueba.mean():.2f} con 0.25 a {ht[ht.temperado == 1.25].logdens_prueba.mean():.2f} con 1.25, y en la espacial es máximo con 0.75 ({he[he.temperado == 0.75].logdens_prueba.mean():.2f}). El valor publicado queda a {mej_t.logdens_prueba - pub_t.logdens_prueba:.2f} del mejor de la prueba temporal y a {mej_e.logdens_prueba - pub_e.logdens_prueba:.2f} del mejor de la espacial: la selección por calibración es razonable, pero no óptima fuera de la zona de entrenamiento.")
add(f"Estabilidad temporal. Con los 2 257 eventos del COE de 2010 a mayo de 2024, la captura anual del 10 % superior del construido está entre {100 * an.captura.min():.1f} % y {100 * an.captura.max():.1f} %, y el límite inferior del intervalo de Wilson (1927) supera 10 % en los {len(an)} años (figura S5b). La correlación de rangos entre año y captura es de −0.46 (p = {pv:.2f}): no hay tendencia significativa.")
add(f"Dependencia espacial. Los eventos del COE presentan una correlación intraclase de {dep.icc.iloc[0]:.2f} entre bloques de 5 km en 2022 a mayo de 2024 y de {dep.icc.iloc[1]:.2f} en 2010 a mayo de 2024, con efectos de diseño de {dep.efecto_diseno.iloc[0]:.1f} y {dep.efecto_diseno.iloc[1]:.1f} (Kish, 1965); el tamaño efectivo es de {dep.n_efectivo.iloc[0]:.0f} de {int(dep.n.iloc[0])} y de {dep.n_efectivo.iloc[1]:.0f} de {int(dep.n.iloc[1])} eventos. Los intervalos por bloques de las tablas del texto incorporan esa dependencia.")
S8.append(cl(cap_tpl, 'Figura S5. Coeficientes del ancla y estabilidad anual de la captura'))
S8.append(picture(f'{FIG}/figS5_posestimacion.png'))
S8.append(cl(note_tpl, 'Fuente: ancla reajustada con bootstrap por bloques de 5 km y eventos de inundación del COE municipal. Elaboración propia. Panel a: intervalos del 95 % de los coeficientes; panel b: intervalos de Wilson del 95 % por año; la línea punteada es el azar (10 %).'))
S8.append(cl(cap_tpl, 'Tabla S9. Resumen de las pruebas posestimación'))
S8.append(make_table([2300, 3900, 2872], ['Prueba', 'Estadístico', 'Lectura'], [
 ['Colinealidad', f'Rango 21 de 22; condición 131 (fondo) y 83 (reportes); VIF > 10 en {vif10} covariables', 'Dos covariables de construido son colineales'],
 ['Estabilidad de coeficientes', f'{n_exc} de 22 con IC sin cero; tres estables en signo (≥ 99.7 %)', 'El desempeño descansa en tres covariables'],
 ['Jackknife por bloque', f'42 bloques; máximo {inf.max_dif_estandarizada.max():.2f} errores estándar; Δ log densidad {inf.dif_logdens_test.min():.3f} a +{inf.dif_logdens_test.max():.3f}', 'Ningún bloque determina el resultado'],
 ['Localización', f'Posición central: {lc.logdens_prueba_25_posiciones:.3f} y {100 * lc.captura_25_posiciones:.1f} %; nueve posiciones: {l9.logdens_prueba_25_posiciones:.3f} y {100 * l9.captura_25_posiciones:.1f} %', 'Resultado estable; evaluar solo en la posición central lo sobrestima'],
 ['Penalización y temperado', f'Log densidad {ht.logdens_prueba.min():.2f} a {ht.logdens_prueba.max():.2f} (temporal) y {he.logdens_prueba.min():.2f} a {he.logdens_prueba.max():.2f} (espacial)', 'El temperado domina; el valor publicado es razonable'],
 ['Estabilidad por año (COE)', f'Captura {100 * an.captura.min():.1f} % a {100 * an.captura.max():.1f} %; IC inferior > 10 % en {len(an)} años', f'Sin tendencia (Spearman −0.46; p = {pv:.2f})'],
 ['Dependencia por bloques', f'ICC {dep.icc.iloc[1]:.2f} a {dep.icc.iloc[0]:.2f}; efecto de diseño {dep.efecto_diseno.iloc[0]:.1f} a {dep.efecto_diseno.iloc[1]:.1f}', 'Justifica el bootstrap por bloques']]))
S8.append(cl(note_tpl, 'Fuente: pruebas posestimación del estudio. Elaboración propia. Detalle en los archivos A8_*.csv de la carpeta de análisis de robustez.'))
for el in S8: ref_head._p.addprevious(el)

# intro del suplemento
p3 = d.paragraphs[3]; r = p3.runs[0]
assert 'y la definición y verificación del NBI (S7). Las tablas' in r.text
r.text = r.text.replace('y la definición y verificación del NBI (S7). Las tablas', 'la definición y verificación del NBI (S7) y las pruebas posestimación (S8). Las tablas').replace('(S6), la definición', '(S6), la definición')

# referencias nuevas (orden alfabético)
NEW = [
 [('Belsley, D. A., Kuh, E., & Welsch, R. E. (1980). Regression diagnostics: Identifying influential data and sources of collinearity. Wiley.', 0)],
 [('Dormann, C. F., Elith, J., Bacher, S., Buchmann, C., Carl, G., Carré, G., García Marquéz, J. R., Gruber, B., Lafourcade, B., Leitão, P. J., Münkemüller, T., McClean, C., Osborne, P. E., Reineking, B., Schröder, B., Skidmore, A. K., Zurell, D., & Lautenbach, S. (2013). Collinearity: A review of methods to deal with it and a simulation study evaluating their performance. ', 0), ('Ecography, 36', 1), ('(1), 27-46. https://doi.org/10.1111/j.1600-0587.2012.07348.x', 0)],
 [('Kish, L. (1965). Survey sampling. Wiley.', 0)],
 [("O'Brien, R. M. (2007). A caution regarding rules of thumb for variance inflation factors. ", 0), ('Quality & Quantity, 41', 1), ('(5), 673-690. https://doi.org/10.1007/s11135-006-9018-6', 0)],
 [('Wilson, E. B. (1927). Probable inference, the law of succession, and statistical inference. ', 0), ('Journal of the American Statistical Association, 22', 1), ('(158), 209-212. https://doi.org/10.1080/01621459.1927.10502953', 0)],
]
i_ref = [i for i, p in enumerate(d.paragraphs) if p.text == 'Referencias'][0]
refs = d.paragraphs[i_ref + 1:]
r0 = [p for p in refs if len(p.runs) == 3][0]            # plantilla con cursiva en el medio
runs_t = [r0.runs[0]._r, r0.runs[1]._r]
have = ' '.join(p.text for p in refs)
for n in ['Künsch', 'Efron']: assert n in have, n
for segs in NEW:
    first = ''.join(t for t, _ in segs)
    new = copy.deepcopy(r0._p)
    for r_ in new.xpath('./w:r'): new.remove(r_)
    for t, it in segs:
        rr = copy.deepcopy(runs_t[it]); new.append(rr)
        Paragraph(new, None).runs[-1].text = t
    # inserción alfabética
    refs = d.paragraphs[i_ref + 1:]
    for p in refs:
        if p.text.casefold() > first.casefold(): p._p.addprevious(new); break
    else: refs[-1]._p.addnext(new)
d.save(SUP)

# ---------------- manuscrito principal ----------------
shutil.copy(MAN, f'{MAN}.bak')
m = docx.Document(MAN)
def app(start, extra):
    p = [p for p in m.paragraphs if p.text.startswith(start)][0]
    assert 'S8' not in p.text
    p.runs[-1].text = p.runs[-1].text + extra
app('donde q es el cuantil 0.9', ' Las pruebas posestimación (colinealidad, influencia por bloques, incertidumbre de localización, penalización y temperado, estabilidad anual y dependencia entre bloques) se reportan en la sección S8.')
app('Cuando la evaluación se restringe al construido', ' Las pruebas posestimación (sección S8) indican que el desempeño descansa en tres covariables (relieve local, construido a 100 m y distancia a vías), que ningún bloque de 5 km lo determina y que la captura anual del 10 % superior con eventos del COE se mantiene entre 16.3 % y 30.0 % en los 15 años, con límites inferiores de Wilson sobre 10 %.')
note = [p for p in m.paragraphs if p.text.startswith('Fuente: evaluación del estudio y eventos de inundación')][0]
assert 'S3, S4 y S6.' in note.text
for r in note.runs:
    if 'S3, S4 y S6.' in r.text: r.text = r.text.replace('S3, S4 y S6.', 'S3, S4, S6 y S8.')
m.save(MAN)
print('ok')
