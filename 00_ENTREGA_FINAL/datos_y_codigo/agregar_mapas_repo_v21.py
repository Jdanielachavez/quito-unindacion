"""Agrega al suplemento (S10) y refiere desde el artículo los mapas del repositorio original."""
import copy, re, docx
from docx.text.paragraph import Paragraph
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT = '/Users/daniela/Desktop/PAPER INUNDACION/'; E = ROOT + '00_ENTREGA_FINAL/'; M = E + 'mapas_repo_original/'
S = docx.Document(E + '2_MATERIAL_SUPLEMENTARIO_V21.docx'); P = S.paragraphs
assert not any(p.text.startswith('S10 ') for p in P)
h2 = [p for p in P if p.text.startswith('S9 Validación')][0]; body = [p for p in P if p.text.startswith('Esta sección contrasta')][0]
cap = [p for p in P if p.text.startswith('Figura S6.')][0]; note = [p for p in P if p.text.startswith('Fuente: Sentinel-1 RTC y eventos del COE')][0]
ref = [p for p in P if p.text == 'Referencias'][0]
def cl(tpl, text):
    new = copy.deepcopy(tpl._p)
    for el in new.xpath('.//w:lastRenderedPageBreak'): el.getparent().remove(el)
    p = Paragraph(new, tpl._parent)
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    p.runs[0].text = text; return new
def pic(path, w):
    new = copy.deepcopy(body._p)
    for r in new.xpath('./w:r'): new.remove(r)
    p = Paragraph(new, body._parent); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(path, width=Cm(w)); return new
els = [cl(h2, 'S10 Mapas del repositorio original'),
 cl(body, 'Las figuras S7 y S8 reproducen, sin modificarlas, los dos mapas que publica el repositorio de replicación del modelo territorial (Morales Lastra, 2026). Las figuras 1 y 3 del artículo se construyeron con los mismos datos, pero la figura 3 usa el NBI por hogar recalculado desde microdatos; por eso la escala de la prioridad con NBI llega a cerca de 19 en el artículo y a cerca de 22 en el mapa original, que usa la tasa parroquial de la capa municipal.'),
 cl(cap, 'Figura S7. Construido analizado y decil superior del score (mapa del repositorio original)'), pic(M + 'FIGURA_construido_prioritario.png', 14),
 cl(note, 'Fuente: repositorio de replicación del modelo (FIGURA_construido_prioritario.png). Panel a: Distrito Metropolitano de Quito; panel b: núcleo urbano y valles. Gris: construido analizado; rojo: decil superior del score.'),
 cl(cap, 'Figura S8. Prioridad por área y prioridad con NBI (mapa del repositorio original)'), pic(M + 'FIGURA_prioridad_integrada.png', 16),
 cl(note, 'Fuente: repositorio de replicación del modelo (FIGURA_prioridad_integrada.png). Prioridad con NBI según la tasa parroquial de la capa municipal; la figura 3 del artículo usa el NBI por hogar recalculado.')]
for e in els: ref._p.addprevious(e)
r = S.paragraphs[3].runs[0]; old = 'la validación con imágenes Sentinel-1 (S9). Las tablas'
assert old in r.text; r.text = r.text.replace(old, 'la validación con imágenes Sentinel-1 (S9) y los mapas del repositorio original (S10). Las tablas')
S.save(E + '2_MATERIAL_SUPLEMENTARIO_V21.docx')
A = docx.Document(E + '1_MANUSCRITO_V21.docx')
n = [p for p in A.paragraphs if p.text.startswith('Fuente: construido prioritario, probabilidades húmedas')]
assert n, 'nota figura 3'
n[0].runs[-1].text = n[0].runs[-1].text.rstrip() + ' Los mapas originales del repositorio se muestran en las figuras S7 y S8.'
A.save(E + '1_MANUSCRITO_V21.docx'); print('ok')
