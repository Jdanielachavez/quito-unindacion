"""Auditoría de la V21 y de su material suplementario: cifras, numeración, referencias cruzadas, citas y formato de resumen.
Uso: python auditoria_v21.py"""
import os, re, unicodedata
import docx
ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
M = docx.Document(f'{ROOT}/01_MANUSCRITO/MANUSCRITO_QUITO_V21.docx'); S = docx.Document(f'{ROOT}/01_MANUSCRITO/MATERIAL_SUPLEMENTARIO_V21.docx')
ok, bad = [], []
def check(n, c, det=''): (ok if c else bad).append(n + (f' -> {det}' if det and not c else ''))
NB = ' '; norm = lambda t: t.replace(NB, ' ')
mt = norm('\n'.join(p.text for p in M.paragraphs)); st = norm('\n'.join(p.text for p in S.paragraphs))
def seg(t, a, b):
    i = t.index(a); return t[i:t.index(b, i)]
# --- formato tipo revista
res = seg(mt, 'Resumen\n', 'Abstract\n'); check('resumen ≤ 250 palabras', len(res.split()) <= 252, str(len(res.split())))
hl = [p.text[2:] for p in M.paragraphs if p.text.startswith('• ')]; check('aspectos destacados: 3 a 5, ≤ 85 caracteres', 3 <= len(hl) <= 5 and all(len(h) <= 85 for h in hl), str([len(h) for h in hl]))
check('palabras clave y JEL presentes', 'Palabras clave:' in mt and 'Códigos JEL' in mt)
check('autores', 'Wilson Gustavo Morales Lastra y Josseline Daniela Chávez Carpio' in mt and 'Josseline Daniela Chávez Carpio' in st)
check('sin nota de «Investigador independiente» ni correo', 'Investigador independiente' not in mt + st and 'wilson.morales@epn' not in mt + st)
for h in ('1 Introducción', '2 Datos y métodos', '3 Resultados', '4 Discusión', '4.2 Limitaciones', '5 Conclusiones', 'Declaraciones', 'Referencias'): check(f'sección «{h}»', h in mt)
body = seg(mt, '1 Introducción\n', 'Declaraciones\n'); print('palabras del cuerpo (1 a 5):', len(body.split()), '| resumen:', len(res.split()), '| suplemento:', len(st.split()))
# --- cifras clave del artículo
for s_ in ('1.48 %', '3.10 %', '0.733', '114.98', '116.68', '1 072.61', '1 039.32', '2 993 619', '367 058', '367 089', '390 464', '23 375', '23 550', '0.444', '0.625', '0.9997', '0.983', '0.998', '9.36 %', '10.55 %', '46 %', '43 %', '23 % y 34 %', '19.1 %', '12.0 a 26.0', '12.4 %', '1.878', '86.8 %', '63.9 %', '91.5 %', '90.4 %', '5.95', '24.96 %', '19.01 %', '28 parroquias'):
    check(f'cifra «{s_}»', s_ in mt)
# --- numeración y referencias cruzadas
def caps(t, kind, pref=''): return [m.group(1) for m in re.finditer(rf'^{kind} ({pref}\d+)\.', t, re.M)]
fm, tm = caps(mt, 'Figura'), caps(mt, 'Tabla'); fs, ts = caps(st, 'Figura', 'S'), caps(st, 'Tabla', 'S')
check('figuras del texto 1..5 en orden', fm == [str(i) for i in range(1, len(fm) + 1)], str(fm)); check('tablas del texto 1..5 en orden', tm == [str(i) for i in range(1, len(tm) + 1)], str(tm))
check('figuras del suplemento S1.. en orden', fs == [f'S{i}' for i in range(1, len(fs) + 1)], str(fs)); check('tablas del suplemento S1.. en orden', ts == [f'S{i}' for i in range(1, len(ts) + 1)], str(ts))
for kind, labs, text in (('figura', fm, mt), ('tabla', tm, mt), ('figura', fs, mt + st), ('tabla', ts, mt + st)):
    for lab in labs: check(f'{kind} {lab} citada', re.search(rf'\b{kind}s? (?:\w+ y )?{lab}\b|\b{kind}s \w+ y {lab}\b', text) is not None)
for m in sorted(set(re.findall(r'\b(?:figura|tabla)s? (S\d+)\b', mt))):
    kind = 'figura' if re.search(rf'figura S{m[1:]}\b', mt) else 'tabla'; check(f'el texto remite a {kind} {m}, que existe en el suplemento', m in (fs if kind == 'figura' else ts))
for m in sorted(set(re.findall(r'\bsecci[oó]nes? (S\d)', mt))): check(f'sección {m} existe en el suplemento', re.search(rf'^{m} ', st, re.M) is not None)
for m in sorted(set(re.findall(r'\bsección (\d\.\d)\b', mt + st))): check(f'sección {m} existe en el texto', re.search(rf'^{re.escape(m)} ', mt, re.M) is not None)
# --- ecuaciones numeradas y citadas
nums = [c.text.strip() for t in M.tables for row in t.rows for c in row.cells if re.fullmatch(r'\(\d\)', c.text.strip())]; check('ocho ecuaciones numeradas (1) a (8)', nums == [f'({i})' for i in range(1, 9)], str(nums))
check('ecuación (2) citada', 'ecuación (2)' in mt)
# --- citas ↔ referencias
def refs(doc_text):
    i = doc_text.index('Referencias\n'); return [l for l in doc_text[i:].split('\n')[1:] if l.strip()], doc_text[:i]
def key(l):
    m = re.match(r'^(.+?)\. \((\d{4}|s\. f\.(?:-[a-z])?)\)', l); 
    if not m: return None
    a = m.group(1).split(',')[0].strip(); return {'Instituto Nacional de Estadística y Censos': 'INEC', 'Servicio Nacional de Gestión de Riesgos': 'SNGR'}.get(a, a), m.group(2)
for nombre, text in (('texto', mt), ('suplemento', st)):
    rl, bd = refs(text)
    check(f'{nombre}: referencias en formato APA', all(key(l) for l in rl))
    for l in rl:
        k = key(l); check(f'{nombre}: «{k[0]} {k[1]}» citada', re.search(re.escape(k[0]) + r'[^;)]{0,70}?' + re.escape(k[1]), bd) is not None)
    for m in sorted(set(re.findall(r'\(([^()]*?(?:\d{4}|s\. f\.(?:-[a-z])?))\)', bd))):
        for piece in re.split(r';\s*', m):
            mm = re.match(r'^([A-ZÁÉÍÓÚÑ][^,&]*?)(?: et al\.| & .+?)?, (\d{4}|s\. f\.(?:-[a-z])?)$', piece.strip())
            if mm:
                a = mm.group(1).split(',')[0].strip(); a = {'Municipio del Distrito Metropolitano de Quito': 'Municipio del Distrito Metropolitano de Quito'}.get(a, a)
                found = any((kk := key(l)) and kk[1] == mm.group(2) and (kk[0].startswith(a.split()[0]) or a.startswith(kk[0])) for l in rl)
                check(f'{nombre}: la cita «{piece.strip()[:50]}» tiene referencia', found)
body_sin_url = re.sub(r'https?://\S+', '', seg(mt, '1 Introducción\n', 'Referencias\n'))
check('sin subíndices con guion bajo en la prosa', not re.findall(r'\b[A-Za-z]_[A-Za-z0-9]+\b', body_sin_url), str(re.findall(r'\b[A-Za-z]_[A-Za-z0-9]+\b', body_sin_url)[:5]))
# --- obsoletos
for s_ in ('21.08', '22.15', '134.56', '1.21 %', '3.51 %', 'sexto', 'cuatro con 20', 'El Condado y Tumbaco ingresan'): check(f'sin obsoleto «{s_}»', s_ not in mt + st)
print(f'OK: {len(ok)} | FALLAN: {len(bad)}'); [print('  ✗', b) for b in bad]
