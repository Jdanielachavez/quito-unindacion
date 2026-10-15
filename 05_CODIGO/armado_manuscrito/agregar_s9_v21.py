"""Agrega la sección S9 (validación con Sentinel-1) al suplemento V21 y sus referidos al manuscrito (edición en su lugar)."""
import copy, re, shutil
import docx, pandas as pd
from docx.text.paragraph import Paragraph
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
ROOT = '/Users/daniela/Desktop/PAPER INUNDACION'
MAN, SUP = f'{ROOT}/01_MANUSCRITO/MANUSCRITO_QUITO_V21.docx', f'{ROOT}/01_MANUSCRITO/MATERIAL_SUPLEMENTARIO_V21.docx'
R09, FIG = f'{ROOT}/09_ANALISIS_ROBUSTEZ', f'{ROOT}/10_FIGURAS'
NB = ' '
def nbs(t): return re.sub(r'(?<=\d) (?=\d{3}(?!\d))', NB, t)
d9 = pd.read_csv(f'{R09}/B2_sentinel1_por_adquisicion.csv'); rs = pd.read_csv(f'{R09}/B2_sentinel1_resumen.csv').set_index('medida')
g = lambda i: rs.iloc[i]
n, npos = len(d9), int((d9.diferencia > 0).sum())
from scipy.stats import binomtest
psig = binomtest(npos, n).pvalue
assert n == 59 and npos == 39

d = docx.Document(SUP); P = d.paragraphs
assert not any(p.text.startswith('S9 ') for p in P)
find = lambda s: [p for p in P if p.text.startswith(s)][0]
h2_tpl, body_tpl = find('S8 Pruebas'), find('Esta sección reúne')
cap_tpl, note_tpl = find('Tabla S9.'), find('Fuente: pruebas posestimación')
tbl_tpl = d.tables[-1]._tbl; ref_head = find('Referencias')
def cl(tpl, text):
    new = copy.deepcopy(tpl._p)
    for el in new.xpath('.//w:lastRenderedPageBreak'): el.getparent().remove(el)
    p = Paragraph(new, tpl._parent)
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    p.runs[0].text = nbs(text); return new
def make_table(widths, header, rows):
    t = copy.deepcopy(tbl_tpl); trs = t.xpath('./w:tr'); hdr, dat = copy.deepcopy(trs[0]), copy.deepcopy(trs[1])
    for tr in trs: t.remove(tr)
    grid = t.find(qn('w:tblGrid'))
    for gg in list(grid): grid.remove(gg)
    for w in widths: etree.SubElement(grid, qn('w:gridCol')).set(qn('w:w'), str(w))
    def build(tpl_tr, vals):
        tr = copy.deepcopy(tpl_tr); tcs = tr.xpath('./w:tc'); tc0 = copy.deepcopy(tcs[0])
        for tc in tcs: tr.remove(tc)
        for w, v in zip(widths, vals):
            tc = copy.deepcopy(tc0); tc.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'), str(w))
            ps = tc.xpath('./w:p')
            for ex in ps[1:]: tc.remove(ex)
            p = Paragraph(ps[0], None)
            for rr in p.runs[1:]: rr._r.getparent().remove(rr._r)
            p.runs[0].text = nbs(str(v)); tr.append(tc)
        return tr
    t.append(build(hdr, header))
    for row in rows: t.append(build(dat, row))
    return t
def picture(path, width_cm=16):
    new = copy.deepcopy(body_tpl._p)
    for r in new.xpath('./w:r'): new.remove(r)
    p = Paragraph(new, body_tpl._parent); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm)); return new

f2 = lambda x: f'{x:.2f}'
S = [cl(h2_tpl, 'S9 Validación con imágenes Sentinel-1')]
add = lambda t: S.append(cl(body_tpl, t))
add('Esta sección contrasta la susceptibilidad con el cambio de retrodispersión de radar de apertura sintética (Sentinel-1, polarización VV, con corrección radiométrica del terreno; Small, 2011) alrededor de eventos de inundación del COE. Se usaron las 59 adquisiciones de 2021 a 2024 que siguen, en 0 a 2 días, a al menos tres eventos de inundación (472 eventos en 384 bloques de 80 m del construido), obtenidas del Microsoft Planetary Computer (Microsoft Open Source et al., 2022). Para cada adquisición, el cambio es la razón en decibelios entre la imagen y la mediana de hasta seis adquisiciones previas de la misma órbita relativa (de 11 a 100 días antes y sin eventos), promediada en bloques de 80 m, que es el nivel de las pirámides del producto y limita el moteado.')
add(f"Los bloques con eventos muestran un cambio de retrodispersión {f2(g(0).media)} dB mayor que el resto del construido (IC 95 %: {f2(g(0).ic95_inf)} a {f2(g(0).ic95_sup)}; {npos} de {n} adquisiciones con diferencia positiva, prueba de signos p = {psig:.2f}) y una proporción de bloques con aumento de más de 3 dB {g(3).media:.1f} puntos porcentuales mayor ({g(3).ic95_inf:.1f} a {g(3).ic95_sup:.1f}). No hay exceso de caídas de más de 3 dB ({g(2).media:.1f} puntos; {g(2).ic95_inf:.1f} a {g(2).ic95_sup:.1f}), que serían la firma del agua abierta. El aumento es la respuesta esperada en áreas urbanas, donde el agua superficial junto a las edificaciones eleva la retrodispersión por doble rebote (Pulvirenti et al., 2016). La correlación de rangos entre el percentil de susceptibilidad y el cambio es de {f2(g(4).media)} ({f2(g(4).ic95_inf)} a {f2(g(4).ic95_sup)}), y los bloques con eventos están {g(5).media / 100:.1f} puntos de percentil por encima del construido promedio ({g(5).ic95_inf / 100:.1f} a {g(5).ic95_sup / 100:.1f}), en línea con la captura de la validación con el COE (tabla 5).")
add('La señal es débil y admite otras explicaciones: los eventos se concentran en zonas más densas, con mayor retrodispersión basal, y la humedad del suelo también eleva la señal. El resultado respalda la dirección de la susceptibilidad, pero no mide extensión ni profundidad, y la resolución de 80 m, el paso de la órbita cada 12 días y la corta duración de la inundación pluvial limitan la detección (Misra et al., 2025; Tellman et al., 2021). Los intervalos usan bootstrap sobre adquisiciones (5 000 réplicas) y, por tanto, no suponen independencia entre bloques de una misma imagen (tabla S10; figura S6).')
S.append(cl(cap_tpl, 'Tabla S10. Cambio de retrodispersión Sentinel-1 en eventos de inundación del COE'))
S.append(make_table([5300, 1300, 1300, 1272], ['Medida', 'Media', 'IC 95 % inf.', 'IC 95 % sup.'], [
 ['Diferencia de cambio VV, eventos menos construido (dB)', f2(g(0).media), f2(g(0).ic95_inf), f2(g(0).ic95_sup)],
 ['Bloques con aumento > 3 dB, eventos menos construido (puntos porcentuales)', f'{g(3).media:.1f}', f'{g(3).ic95_inf:.1f}', f'{g(3).ic95_sup:.1f}'],
 ['Bloques con caída > 3 dB, eventos menos construido (puntos porcentuales)', f'{g(2).media:.1f}', f'{g(2).ic95_inf:.1f}', f'{g(2).ic95_sup:.1f}'],
 ['Correlación de rangos entre susceptibilidad y cambio VV', f2(g(4).media), f2(g(4).ic95_inf), f2(g(4).ic95_sup)],
 ['Percentil de susceptibilidad, eventos menos construido (puntos)', f'{g(5).media / 100:.1f}', f'{g(5).ic95_inf / 100:.1f}', f'{g(5).ic95_sup / 100:.1f}']]))
S.append(cl(note_tpl, f'Fuente: Sentinel-1 RTC (Microsoft Planetary Computer) y eventos de inundación del COE municipal. Elaboración propia. {n} adquisiciones de 2021 a 2024; 472 eventos; bloques de 80 m del construido. Detalle en B2_sentinel1_*.csv.'))
S.append(cl(cap_tpl, 'Figura S6. Diferencia de cambio de retrodispersión entre bloques con eventos y el resto del construido'))
S.append(picture(f'{FIG}/figS6_sentinel1.png'))
S.append(cl(note_tpl, 'Fuente: Sentinel-1 RTC y eventos del COE municipal. Elaboración propia. Panel a: cada punto es una adquisición (tamaño proporcional al número de bloques con eventos); la línea discontinua y la banda son la media y su IC 95 %. Panel b: correlación de rangos por adquisición; la línea discontinua es la media.'))
for el in S: ref_head._p.addprevious(el)

# intro
r = d.paragraphs[3].runs[0]
old = 'y las pruebas posestimación (S8). Las tablas'
assert old in r.text; r.text = r.text.replace(old, 'las pruebas posestimación (S8) y la validación con imágenes Sentinel-1 (S9). Las tablas').replace('la definición y verificación del NBI (S7) las pruebas', 'la definición y verificación del NBI (S7), las pruebas')
print(d.paragraphs[3].text)

# referencias
NEW = [
 [('Microsoft Open Source, McFarland, M., Emanuele, R., Morris, D., & Augspurger, T. (2022). ', 0), ('microsoft/PlanetaryComputer', 1), (': October 2022 [Software]. Zenodo. https://doi.org/10.5281/zenodo.7261897', 0)],
 [('Pulvirenti, L., Chini, M., Pierdicca, N., & Boni, G. (2016). Use of SAR data for detecting floodwater in urban and agricultural areas: The role of the interferometric coherence. ', 0), ('IEEE Transactions on Geoscience and Remote Sensing, 54', 1), ('(3), 1532-1544. https://doi.org/10.1109/TGRS.2015.2482001', 0)],
 [('Small, D. (2011). Flattening gamma: Radiometric terrain correction for SAR imagery. ', 0), ('IEEE Transactions on Geoscience and Remote Sensing, 49', 1), ('(8), 3081-3093. https://doi.org/10.1109/TGRS.2011.2120616', 0)],
]
i_ref = [i for i, p in enumerate(d.paragraphs) if p.text == 'Referencias'][0]
r0 = [p for p in d.paragraphs[i_ref + 1:] if len(p.runs) == 3][0]
tpls = [r0.runs[0]._r, r0.runs[1]._r]
for segs in NEW:
    first = ''.join(t for t, _ in segs)
    new = copy.deepcopy(r0._p)
    for x in new.xpath('./w:r'): new.remove(x)
    for t, it in segs:
        rr = copy.deepcopy(tpls[it]); new.append(rr); Paragraph(new, None).runs[-1].text = t
    for p in d.paragraphs[i_ref + 1:]:
        if p.text.casefold() > first.casefold(): p._p.addprevious(new); break
    else: d.paragraphs[-1]._p.addnext(new)
d.save(SUP)

# manuscrito principal
m = docx.Document(MAN)
def app(start, old, new):
    p = [p for p in m.paragraphs if p.text.startswith(start)][0]
    for r in p.runs:
        if old in r.text: r.text = r.text.replace(old, new, 1); return
    raise SystemExit('sin coincidencia: ' + start)
app('Cuando la evaluación se restringe al construido', 'con límites inferiores de Wilson sobre 10 %.',
    'con límites inferiores de Wilson sobre 10 %. Con 59 imágenes Sentinel-1 de 2021 a 2024, los bloques con eventos muestran un aumento de retrodispersión 0.25 dB mayor que el resto del construido (IC 95 %: 0.10 a 0.40) y la susceptibilidad se asocia con ese cambio (correlación de rangos de 0.07), una señal débil y de la dirección esperada para agua urbana (sección S9).')
app('Primero, la muestra de reportes', 'Por último, el estudio no estima extensión inundada, profundidad, daños ni personas afectadas.',
    'Por último, el estudio no estima extensión inundada, profundidad, daños ni personas afectadas; la validación con radar (sección S9) es un contraste de señal con resolución de 80 m y no una cartografía de inundación.')
app('La capacidad del modelo para discriminar dentro del construido es modesta', 'validar con fuentes de eventos independientes y con observaciones físicas, medir la privación a escala de sector y vincular',
    'validar con registros de eventos con fecha y dirección y con cartografía de inundación de mayor resolución que la de Sentinel-1, medir la privación en hogares dentro de cada sector censal y vincular')
note = [p for p in m.paragraphs if p.text.startswith('Fuente: evaluación del estudio y eventos de inundación')][0]
for r in note.runs:
    if 'S3, S4, S6 y S8.' in r.text: r.text = r.text.replace('S3, S4, S6 y S8.', 'S3, S4, S6, S8 y S9.')
m.save(MAN); print('ok')
